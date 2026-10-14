# -*- coding: utf-8 -*-
"""
md_loader.py —— markdown 解析（复刻 md_loader.py）

生产对照（aichain/document_loader/md_loader.py:15-33）：
    UnstructuredPaddleMDLoader._get_elements(:18-22)
        pandoc -s -o <docx> <md> 转成 docx → partition_docx（标题分块）

demo 与生产的差异：
    · 无 pandoc 的纯 Python fallback（_md_to_docx_by_python：# → add_heading，
      表格块 → 真表格，保留"md→docx→标题分块"的管线本质）；生产环境恒有
      pandoc，直接 subprocess.run
    · _markdown_text_to_docx 同时是 pdf 路线（MinerU → markdown）的共用转换，
      结构映射：# 标题→Heading、|表格|→docx 表格、图片行跳过、其余→段落
"""
import os
import re
import shutil
import subprocess
from typing import Any, List

import docx as docx_lib
from langchain_core.documents import Document
from langchain_core.document_loaders import BaseLoader

from .word_loader import partition_docx


def _markdown_text_to_docx(md_text: str, docx_path: str):
    """markdown 文本 → docx —— pdf(MinerU) 与 md 两条路线共用的转换。

    结构映射（对照生产解析产物，保证后续 partition_docx 同路处理）：
      · `#`~`######` 标题 → add_heading(level)（标题合并/父子切分依赖它）
      · markdown 表格块（|表头| + |---| 分隔行 + 数据行）→ docx 真表格
        （add_table；partition_docx 侧再按生产 _convert_table_to_text 拉平）
      · 图片行 `![alt](url/base64)` → 跳过（对向量检索是噪声；生产图片走
        minio 提取，demo 不落盘）
      · 其余行（列表/公式/正文）→ add_paragraph 原文
    """
    d = docx_lib.Document()
    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        text = lines[i].rstrip()
        if not text.strip():
            i += 1
            continue
        # 图片行：跳过（见 docstring）
        if re.match(r'^\s*!\[[^\]]*\]\([^)]*\)\s*$', text):
            i += 1
            continue
        # 标题行
        m = re.match(r'^(#{1,6})\s+(.*)$', text)
        if m:
            d.add_heading(m.group(2), level=len(m.group(1)))
            i += 1
            continue
        # 表格块：|表头| 行 + |---| 分隔行 起头，连续 | 行为一个表
        if text.strip().startswith('|') and i + 1 < len(lines) and \
                re.match(r'^\|[\s:|\-]+\|?$', lines[i + 1].strip()):
            header = [c.strip() for c in text.strip().strip('|').split('|')]
            rows = []
            j = i + 2                                  # 跳过分隔行
            while j < len(lines) and lines[j].strip().startswith('|'):
                rows.append([c.strip() for c in lines[j].strip().strip('|').split('|')])
                j += 1
            tbl = d.add_table(rows=1 + len(rows), cols=len(header))
            try:
                tbl.style = 'Table Grid'
            except Exception:
                pass                                   # 默认模板缺样式时裸表也可
            for c, val in enumerate(header):
                tbl.rows[0].cells[c].text = val
            for r, row in enumerate(rows, start=1):
                for c in range(len(header)):
                    tbl.rows[r].cells[c].text = row[c] if c < len(row) else ''
            i = j
            continue
        if text.strip() in ("---", "***"):
            i += 1
            continue                                   # 分隔线跳过
        d.add_paragraph(text.strip())
        i += 1
    d.save(docx_path)


def _md_to_docx_by_python(md_path: str, docx_path: str):
    """无 pandoc 时的 fallback：读 md 文件 → _markdown_text_to_docx。

    生产环境恒有 pandoc（md_loader.py:20 直接 subprocess.run），demo 机器可能
    没有，故保留"md→docx→partition_docx"管线本质的同时用纯 Python 转换。"""
    with open(md_path, "r", encoding="utf-8") as f:
        _markdown_text_to_docx(f.read(), docx_path)


class MarkdownDocumentLoader(BaseLoader):
    """md loader —— 复刻 UnstructuredPaddleMDLoader._get_elements(md_loader.py:18-22)：
    pandoc -s -o <docx> <md> 转成 docx 后走 partition_docx；
    demo 加了无 pandoc 的纯 Python fallback（见 _md_to_docx_by_python）。"""

    def __init__(self, file_path: str, mode: str = "elements", **kwargs: Any):
        self.file_path = file_path
        self.mode = mode

    def load(self) -> List[Document]:
        print(f"[loader] md: {os.path.basename(self.file_path)}", flush=True)
        docs_path = self.file_path.replace('.md', '.docx')
        if shutil.which("pandoc"):
            # 生产 md_loader.py:20 原样
            subprocess.run(["pandoc", "-s", "-o", docs_path, self.file_path], check=True)
        else:
            print("    [loader] 未安装 pandoc，用纯 Python md→docx fallback"
                  "（生产用 pandoc）", flush=True)
            _md_to_docx_by_python(self.file_path, docs_path)
        docs = partition_docx(docs_path)
        for d in docs:
            d.metadata["source"] = os.path.basename(self.file_path)
        return docs
