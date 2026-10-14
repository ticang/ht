# -*- coding: utf-8 -*-
"""
word_loader.py —— word（docx/doc）解析（复刻 word_partition.py + word_document.py）

生产对照（aichain/document_loader/word_partition.py）：
    _get_title_level(:458)                    样式名判定标题级别（DOC_STYLE_MAPPING）
    _paragraph_to_element(:471)               标题编号剥离正则
    _get_paragraph_possible_font_size(:489)   段落主字号（run 优先/样式回退）
    _is_possible_tile_by_text(:588)           无样式 docx 的标题文本启发
    _is_possible_title(:637)                  字号+文本形态的标题判定
    _convert_table_to_text(:542-563)          表格拉平为文本
    partition_docx(:399-428)                  标题分块聚合（subject=「父标题,子标题」）
    WordDocumentLoader(word_document.py:17)   langchain BaseLoader 包装

demo 与生产的差异（逐处有注释）：
    · 单遍聚合（生产 handle_element 先产 (element, current_title, parent_title)
      元组流再由 partition_docx 聚合，两遍）——实现不同、结果同构
    · 字号启发保留判定三要素（短文本/非句尾标点/字号>正文+0.5），
      删 between_text 等启发细节；正文主字号取全局众数
    · 图片提取转 minio（:107 的 knowledge_id bucket）未复刻，形参保留
"""
import os
import re
from typing import Any, List, Optional, Tuple

import docx as docx_lib
from langchain_core.documents import Document
from langchain_core.document_loaders import BaseLoader

# 生产 word_partition.py:44 OPSGPT_TITLE_SPLIT：subject 的父/子标题分隔符
TITLE_SPLIT = ","


def _strip_numbering(text: str) -> str:
    """剥标题编号（word_partition.py:471 `_paragraph_to_element` 同款正则：
    '1.2 丢弃定时器' → '丢弃定时器'）。"""
    return re.sub(r'^\d+(?:\.\d+){1,4}', '', text).strip()


def _heading_level_by_style(paragraph) -> Optional[int]:
    """样式名判定标题级别（生产 _get_title_level:458 + DOC_STYLE_MAPPING）：
    'Title'→1，'Heading N'→N，其余 None（正文）。"""
    style_name = paragraph.style.name if paragraph.style else ""
    if style_name == "Title":
        return 1
    if style_name.startswith("Heading"):
        m = re.search(r'(\d+)', style_name)
        return int(m.group(1)) if m else 1
    return None


def _paragraph_font_size(paragraph) -> Optional[float]:
    """段落主字号 —— 生产 _get_paragraph_possible_font_size(:489)：取 runs 里
    非空 run 的字号（显式 run 字号优先，无则回退样式字号）。"""
    sizes = [r.font.size.pt for r in paragraph.runs
             if r.text.strip() and r.font.size is not None]
    if sizes:
        return max(set(sizes), key=sizes.count)        # 众数：正文混排时取主流字号
    if paragraph.style and paragraph.style.font and paragraph.style.font.size:
        return paragraph.style.font.size.pt
    return None


def _is_possible_title(text: str, font_size: Optional[float],
                       body_size: Optional[float]) -> bool:
    """无 Heading 样式时的标题启发 —— 生产 _is_possible_title(:637) /
    _is_possible_tile_by_text(:588) 的精简版（pdf2docx 转出的 docx 没有样式，
    生产正是靠这组字号+文本形态启发识别标题）：
      ① 短文本（生产阈值 24 字符）
      ② 不以句尾标点结束（'。；，'等结尾的是句子）
      ③ 字号明显大于正文主字号（>body+0.5）"""
    if not text or len(text) > 24:
        return False
    if text[-1] in "。；;，,．.":
        return False
    if font_size and body_size and font_size > body_size + 0.5:
        return True
    return False


def _body_font_size(document) -> Optional[float]:
    """全文正文主字号：所有段落字号的众数（生产按逐段与前段的比较判定，
    demo 用全局众数，效果等价且更稳）。"""
    sizes = []
    for p in document.paragraphs:
        s = _paragraph_font_size(p)
        if s:
            sizes.append(s)
    if not sizes:
        return None
    return max(set(sizes), key=sizes.count)


def _convert_table_to_text(table) -> str:
    """docx 表格 → 纯文本 —— 复刻生产 _convert_table_to_text(word_partition.py:542-563)：
    每行单元格拼接（内部换行剥掉）、空行跳过（_is_empty_row:537-540）、行间换行。
    表格来源：MinerU markdown 表格转成的 docx 表格、word 文档原生表格。"""
    lines = []
    for row in table.rows:
        cells = [c.text.replace("\n", "").strip() for c in row.cells]
        if any(cells):                                # 空行跳过（生产 _is_empty_row）
            lines.append("".join(cells))
    return "\n".join(lines)


def partition_docx(filename: str, knowledge_id: str = None) -> List[Document]:
    """docx → List[Document]，标题分块聚合。

    生产结构（word_partition.py）：handle_element 先产出
    (element, current_title, parent_title) 元组流（每个正文段落都带它所属的
    当前/父标题引用），partition_docx(:399-428) 再把相邻正文聚合、把
    subject 写进 metadata（父≠子时 '父标题,子标题'，:406-409）。

    demo 用单遍聚合等价实现（原 parse_docx 升级版）：
      · 按文档流顺序遍历 body（段落 + 表格混排都按原顺序处理——生产
        handle_element 同样按 body 顺序；表格转文本走 _convert_table_to_text，
        复刻 :542-563）
      · 标题判定 = Heading/Title 样式 ∨ 字号启发（供 pdf2docx 转出的无样式 docx）
      · level≤1 或首个标题 → parent_title；level≥2 → current_title
      · 正文聚合到当前标题下，flush 时写 subject
    knowledge_id：生产用于图片上传 minio 的 bucket（:107），demo 保留形参不使用。
    """
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx_lib.Document(filename)
    body_size = _body_font_size(document)
    sections: List[Tuple[str, List[str]]] = []          # (subject, [正文段])
    parent_title, current_title, current_body = None, None, []

    def flush():
        if current_title is not None and current_body:
            subject = current_title
            if parent_title and parent_title != current_title:
                subject = TITLE_SPLIT.join([parent_title, current_title])
            sections.append((subject, current_body))
        elif current_body:
            # 兜底：全文无标题（MinerU 的 markdown 无 #、word 无 Heading 且
            # 字号启发未命中）时正文不能丢——用文件名当 subject，等价生产
            # unstructured 对无标题文档的 Title 判定兜底
            sections.append((os.path.splitext(os.path.basename(filename))[0],
                             current_body))

    # 文档流顺序遍历（w:p=段落，w:tbl=表格），保持原始混排顺序
    for child in document.element.body.iterchildren():
        tag = child.tag
        if tag.endswith('}p'):
            para = Paragraph(child, document)
            text = para.text.strip()
            if not text:
                continue
            level = _heading_level_by_style(para)
            if level is None and _is_possible_title(text, _paragraph_font_size(para), body_size):
                level = 1                                # 字号启发命中按一级标题
            if level is not None:
                flush()
                current_body = []
                current_title = _strip_numbering(text) or text
                if level <= 1 or parent_title is None:
                    parent_title = current_title
            else:
                current_body.append(text)
        elif tag.endswith('}tbl'):
            table = Table(child, document)
            table_text = _convert_table_to_text(table)
            if table_text:
                current_body.append(table_text)         # 表格文本并入当前标题正文
    flush()

    # 生产 load()（word_document.py:64-94）：metadata 带 subject/source；
    # 文件名正则 `_后取知识名` 拼 subject 前缀是上传路径约定，demo 用 basename
    return [Document(page_content="\n".join(body),
                     metadata={"subject": subject, "source": os.path.basename(filename)})
            for subject, body in sections]


class WordDocumentLoader(BaseLoader):
    """word loader —— 复刻 WordDocumentLoader(word_document.py:17)。

    生产 mode='elements'（FILE_LOADER_MAP docx 项的 kwargs）：每个聚合后的
    "标题节"一个 Document；load_and_split 用基类默认实现（text_splitter.
    split_documents(self.load())），与生产覆写版(word_document.py:96-105)一致。
    partition_type：生产 !='default' 时动态 import third_part 插件(:53-55)，
    demo 无插件体系，仅保留参数与日志。"""

    def __init__(self, file_path: str, mode: str = "elements",
                 partition_type: str = "default", knowledge_id: str = None, **kwargs: Any):
        self.file_path = file_path
        self.mode = mode
        self.partition_type = partition_type
        self.knowledge_id = knowledge_id

    def load(self) -> List[Document]:
        print(f"[loader] word: {os.path.basename(self.file_path)}"
              f"（partition_type={self.partition_type}）", flush=True)
        return partition_docx(self.file_path, self.knowledge_id)
